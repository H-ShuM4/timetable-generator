"""事務局の様式どおりに Excel を組めているかを見る。

**見本は「2026時間割（新経営学科）.xlsx」である。** 曜日を横に並べ、時限と
年次で縦に区切る。画面の 5×5 マトリクスとは別物なので、ここで確かめるのは
「どのセルに何が入るか」ではなく「どの列・どの行に入るか」になる。

列の位置は学科で変わる（短大だけ開講期の欄がある、大学だけ金曜に教室が
無い）ため、座標を直書きせず、見出しから列を引いて確かめる。そうしないと
列を 1 つ足しただけでテストが総崩れになる。
"""
import openpyxl
import pytest

from app.constraints.context import Context
from app.export.excel_writer import SHEET_PLAN, write_timetable_excel
from app.models.enums import Category, Department, Quarter, Term
from tests.factories import subject as make
from app.models.timeslot import TimeSlot
from app.models.timetable import AssignmentSource, Timetable
from app.scheduler.pipeline import GenerationResult

HEADER_ROW = 3
FIRST_BODY_ROW = 4


def build(subjects, placements, intensive_codes=()):
    ctx = Context.from_lists(subjects, [])
    tt = Timetable()
    for code, slots in placements.items():
        tt.place(code, slots, AssignmentSource.GEMINI)
    return ctx, GenerationResult(timetable=tt, intensive_codes=list(intensive_codes))


def write(tmp_path, subjects, placements, intensive_codes=()):
    ctx, result = build(subjects, placements, intensive_codes)
    return openpyxl.load_workbook(
        write_timetable_excel(ctx, result, tmp_path / "時間割.xlsx")
    )


def day_block(sheet, day: str) -> int:
    """その曜日の見出しが始まる列を返す。"""
    for cell in sheet[2]:
        if cell.value and str(cell.value).startswith(day):
            return cell.column
    raise AssertionError(f"{day}曜日の見出しが無い")


def column_of(sheet, day: str, title: str) -> int:
    """曜日ブロックの中から、その見出しを持つ列を返す。"""
    start = day_block(sheet, day)
    for column in range(start, start + 8):
        if sheet.cell(HEADER_ROW, column).value == title:
            return column
    raise AssertionError(f"{day}曜日に「{title}」の列が無い")


def intensive_block(sheet) -> tuple[int, int]:
    """集中の欄の（年次列, 授業科目名列）を返す。

    見本と同じく、年次の列には見出しを置かない（AG2:AG3 が空欄で、
    「集中」は 1 つ右の AH2 から始まる）。見出しの位置から 1 つ左が年次列。
    """
    heading = next(c.column for c in sheet[2] if c.value == "集中")
    name = next(column for column in range(heading, heading + 6)
                if sheet.cell(HEADER_ROW, column).value == "授業科目名")
    return heading - 1, name


def find_row(sheet, column: int, value: str) -> int:
    for row in range(FIRST_BODY_ROW, sheet.max_row + 1):
        if sheet.cell(row, column).value == value:
            return row
    raise AssertionError(f"{value} が {column} 列に無い")


# ---------------------------------------------------------------- 全体の形

def test_the_workbook_has_one_sheet_per_department_and_term(tmp_path):
    """学科 × 学期の 6 シート。集中は各シートの右端に入るので別シートは持たない。"""
    book = write(tmp_path, [make("A1")], {"A1": (TimeSlot("月", 1),)})
    assert len(SHEET_PLAN) == 6
    assert book.sheetnames == [name for name, _, _ in SHEET_PLAN]


def test_each_sheet_says_which_department_and_term_it_is(tmp_path):
    book = write(tmp_path, [make("A1")], {"A1": (TimeSlot("月", 1),)})
    assert book["経営・前期"]["A1"].value == "経営学科　前期"
    assert book["短大・後期"]["A1"].value == "短期大学部　後期"


def test_the_days_run_across_and_friday_is_marked_remote(tmp_path):
    """H8 により遠隔科目は金曜へ集まる。見本はそれを見出しで示している。"""
    sheet = write(tmp_path, [make("A1")], {"A1": (TimeSlot("月", 1),)})["経営・前期"]
    headings = [c.value for c in sheet[2] if c.value]
    assert headings == [
        "月曜日：対面", "火曜日：対面", "水曜日：対面",
        "木曜日：対面", "金曜日：遠隔", "集中",
    ]


def test_periods_and_years_run_down_the_left(tmp_path):
    sheet = write(tmp_path, [make("A1")], {"A1": (TimeSlot("月", 1),)})["経営・前期"]
    periods = [c.value for c in sheet["A"] if c.value and "時限" in str(c.value)]
    assert periods == ["1時限", "2時限", "3時限", "4時限", "5時限"]
    years = [c.value for c in sheet["B"] if c.value]
    assert years[:4] == ["1年", "2年", "3年", "4年"], "大学は 4 年まで"


def test_the_junior_college_only_goes_up_to_the_second_year(tmp_path):
    sheet = write(tmp_path, [make("J1", department=Department.JUNIOR)],
                  {"J1": (TimeSlot("月", 1),)})["短大・前期"]
    assert {c.value for c in sheet["B"] if c.value} == {"1年", "2年"}


# ---------------------------------------------------------------- 科目の位置

def test_a_subject_lands_in_its_day_period_and_year(tmp_path):
    subject = make("A1", name="経営学入門", teacher="築雅之", year=3)
    sheet = write(tmp_path, [subject], {"A1": (TimeSlot("水", 3),)})["経営・前期"]

    row = find_row(sheet, column_of(sheet, "水", "授業科目名"), "経営学入門")
    assert sheet.cell(row, column_of(sheet, "水", "教員")).value == "築雅之"
    assert sheet.cell(row, 1).value == "3時限" or sheet.cell(row, 1).value is None
    # 時限と年次は縦結合されるので、見出しは各ブロックの先頭行にだけ入る
    assert _merged_label(sheet, row, 1) == "3時限"
    assert _merged_label(sheet, row, 2) == "3年"


def _merged_label(sheet, row: int, column: int) -> str:
    """縦結合された見出しの値を引く。結合の先頭以外は空になる。"""
    for rng in sheet.merged_cells.ranges:
        if rng.min_col == column and rng.min_row <= row <= rng.max_row:
            return sheet.cell(rng.min_row, column).value
    return sheet.cell(row, column).value


def test_a_double_slot_subject_appears_in_both_periods(tmp_path):
    subject = make("J1", name="動画制作", department=Department.JUNIOR,
                   slots_required=2, requires_consecutive=True)
    sheet = write(tmp_path, [subject],
                  {"J1": (TimeSlot("水", 2), TimeSlot("水", 3))})["短大・前期"]

    column = column_of(sheet, "水", "授業科目名")
    rows = [r for r in range(FIRST_BODY_ROW, sheet.max_row + 1)
            if sheet.cell(r, column).value == "動画制作"]
    assert len(rows) == 2, "▲科目は 2 コマとも出る"
    assert _merged_label(sheet, rows[0], 1) == "2時限"
    assert _merged_label(sheet, rows[1], 1) == "3時限"


def test_a_fall_subject_stays_off_the_spring_sheet(tmp_path):
    subject = make("A1", name="後期の科目", term=Term.FALL)
    book = write(tmp_path, [subject], {"A1": (TimeSlot("月", 1),)})

    column = column_of(book["経営・前期"], "月", "授業科目名")
    assert all(book["経営・前期"].cell(r, column).value != "後期の科目"
               for r in range(FIRST_BODY_ROW, book["経営・前期"].max_row + 1))
    find_row(book["経営・後期"], column_of(book["経営・後期"], "月", "授業科目名"),
             "後期の科目")


# ---------------------------------------------------------------- 学科で違う列

def test_the_university_sheets_follow_the_sample(tmp_path):
    """見本どおり、備考には必修の「必」だけを出す（選択必修は空欄）。"""
    required = make("A1", name="必修の科目", year=1, category=Category.REQUIRED)
    chosen = make("A2", name="選択必修の科目", year=2,
                  category=Category.ELECTIVE_REQUIRED, teacher="教員乙")
    sheet = write(tmp_path, [required, chosen], {
        "A1": (TimeSlot("月", 1),), "A2": (TimeSlot("月", 1),),
    })["経営・前期"]

    note = column_of(sheet, "月", "備考")
    name = column_of(sheet, "月", "授業科目名")
    assert sheet.cell(find_row(sheet, name, "必修の科目"), note).value == "必"
    assert sheet.cell(find_row(sheet, name, "選択必修の科目"), note).value is None


def test_the_university_sheets_drop_the_room_on_friday(tmp_path):
    """大学のカリキュラム一覧は遠隔列が ○ か × で埋まっており、空欄が無い。

    H8 により金曜へ来るのは遠隔科目だけになるので、教室の欄が要らない。
    """
    sheet = write(tmp_path, [make("A1")], {"A1": (TimeSlot("月", 1),)})["経営・前期"]
    for day in ("月", "火", "水", "木"):
        column_of(sheet, day, "教室")
    with pytest.raises(AssertionError):
        column_of(sheet, "金", "教室")


def test_the_junior_college_keeps_the_room_on_friday(tmp_path):
    """短大の遠隔列は ○ か空欄で × が無い。対面科目が金曜に来うる。"""
    sheet = write(tmp_path, [make("J1", department=Department.JUNIOR)],
                  {"J1": (TimeSlot("月", 1),)})["短大・前期"]
    for day in ("月", "火", "水", "木", "金"):
        column_of(sheet, day, "教室")


def test_the_junior_college_shows_the_quarter_and_tells_the_categories_apart(tmp_path):
    """短大にはクオーター科目が 47 件あり、大学には 1 件も無い。

    開講期の欄が要るのは短大だけなので、そこだけ科目名の右に「備考」を足し、
    末尾を「必修」に変えて必・選必を書き分ける。
    """
    early = make("J1", name="スタディスキルゼミナールⅠ", department=Department.JUNIOR,
                 quarter=Quarter.Q1, category=Category.REQUIRED)
    late = make("J2", name="ホテルビジネス実務論", department=Department.JUNIOR,
                quarter=Quarter.Q2, category=Category.ELECTIVE_REQUIRED,
                teacher="教員乙")
    plain = make("J3", name="通しの科目", department=Department.JUNIOR,
                 category=Category.ELECTIVE, teacher="教員丙")
    sheet = write(tmp_path, [early, late, plain], {
        "J1": (TimeSlot("火", 2),), "J2": (TimeSlot("火", 2),), "J3": (TimeSlot("火", 2),),
    })["短大・前期"]

    name = column_of(sheet, "火", "授業科目名")
    quarter = column_of(sheet, "火", "備考")
    category = column_of(sheet, "火", "必修")
    assert quarter == name + 1, "開講期は科目名のすぐ右"

    for subject_name, shown_quarter, mark in (
        ("スタディスキルゼミナールⅠ", "前①", "必"),
        ("ホテルビジネス実務論", "前②", "選必"),
        ("通しの科目", None, None),
    ):
        row = find_row(sheet, name, subject_name)
        assert sheet.cell(row, quarter).value == shown_quarter
        assert sheet.cell(row, category).value == mark


# -------------------------------------------------- 同じ科目を複数の教員が持つ

def test_one_subject_with_many_teachers_is_written_once(tmp_path):
    """課題研究Ⅰは 18 名、卒業研究Ⅰは 16 名が担当する。

    科目名を毎行繰り返すと読めなくなるので、見本と同じく名前は 1 回だけ
    書き、教員を続く行に並べて、名前と備考のセルを縦に結合する。
    """
    seminar = [make(f"A{i}", name="課題研究Ⅰ", year=3, teacher=f"教員{i}")
               for i in range(1, 5)]
    sheet = write(tmp_path, seminar,
                  {s.code: (TimeSlot("月", 3),) for s in seminar})["経営・前期"]

    name = column_of(sheet, "月", "授業科目名")
    note = column_of(sheet, "月", "備考")
    teacher = column_of(sheet, "月", "教員")
    row = find_row(sheet, name, "課題研究Ⅰ")

    assert [sheet.cell(row + i, teacher).value for i in range(4)] == [
        "教員1", "教員2", "教員3", "教員4"
    ]
    assert sheet.cell(row + 1, name).value is None, "科目名は 1 回だけ"
    merged = {str(r) for r in sheet.merged_cells.ranges}
    letter = sheet.cell(row, name).column_letter
    mark = sheet.cell(row, note).column_letter
    assert f"{letter}{row}:{letter}{row + 3}" in merged
    assert f"{mark}{row}:{mark}{row + 3}" in merged


def test_the_same_name_in_different_quarters_stays_apart(tmp_path):
    """前①の回と前②の回は別の授業。まとめると開講期が 1 つしか書けない。"""
    first = make("J1", name="スタディスキル", department=Department.JUNIOR,
                 quarter=Quarter.Q1)
    second = make("J2", name="スタディスキル", department=Department.JUNIOR,
                  quarter=Quarter.Q2, teacher="教員乙")
    sheet = write(tmp_path, [first, second], {
        "J1": (TimeSlot("月", 1),), "J2": (TimeSlot("月", 1),),
    })["短大・前期"]

    name = column_of(sheet, "月", "授業科目名")
    quarter = column_of(sheet, "月", "備考")
    rows = [r for r in range(FIRST_BODY_ROW, sheet.max_row + 1)
            if sheet.cell(r, name).value == "スタディスキル"]
    assert len(rows) == 2
    assert {sheet.cell(r, quarter).value for r in rows} == {"前①", "前②"}


# ---------------------------------------------------------------- 集中講義

def test_intensive_subjects_sit_in_the_intensive_column(tmp_path):
    """学科 × 学期を見分けて、その紙の集中の欄へ入れる。"""
    mine = make("A9", name="認定ＰＢＬ", is_intensive=True, year=2)
    other = make("B9", name="よその学科の集中", is_intensive=True,
                 department=Department.ACCOUNTING)
    sheet = write(tmp_path, [mine, other], {}, intensive_codes=["A9", "B9"])["経営・前期"]

    year_column, name = intensive_block(sheet)
    shown = [sheet.cell(r, name).value for r in range(FIRST_BODY_ROW, sheet.max_row + 1)]
    assert "認定ＰＢＬ" in shown
    assert "よその学科の集中" not in shown
    assert _merged_label(sheet, find_row(sheet, name, "認定ＰＢＬ"), year_column) == "2年"


def test_a_full_year_intensive_shows_on_both_terms(tmp_path):
    """通年の集中講義は前期・後期のどちらの紙にも出す。

    片方にだけ載せると、もう片方を見た人が見落とす。実データでは
    模擬ブライダルプロジェクト（短大 1 年）が 1 件ある。
    """
    subject = make("J9", name="模擬ブライダルプロジェクト", is_intensive=True,
                   department=Department.JUNIOR, term=Term.FULL_YEAR)
    book = write(tmp_path, [subject], {}, intensive_codes=["J9"])

    for name in ("短大・前期", "短大・後期"):
        sheet = book[name]
        _, column = intensive_block(sheet)
        find_row(sheet, column, "模擬ブライダルプロジェクト")


def test_the_sheet_is_set_up_for_a3_landscape(tmp_path):
    """見本と同じ A3 横。倍率は決め打ちにせず、幅に合わせる。"""
    sheet = write(tmp_path, [make("A1")], {"A1": (TimeSlot("月", 1),)})["経営・前期"]
    assert sheet.page_setup.orientation == "landscape"
    assert int(sheet.page_setup.paperSize) == int(sheet.PAPERSIZE_A3)
    assert sheet.page_setup.fitToWidth == 1
    assert sheet.print_title_rows == "$1:$3", "見出しを各ページに繰り返す"
    assert sheet.freeze_panes == "C4"


# ---------------------------------------------------------------- 氏名の表記

def test_the_teacher_name_keeps_the_space_the_office_types(tmp_path):
    """突合は空白を落とした形で行うが、紙に出すのは元の表記に戻す。

    カリキュラム一覧の教員氏名は 97% が「築　雅之」のように全角空白で
    姓名を分けており、事務局は長年その形で時間割を作ってきた。
    """
    subject = make("A1", teacher="築雅之", teacher_display="築　雅之")
    sheet = write(tmp_path, [subject], {"A1": (TimeSlot("月", 1),)})["経営・前期"]

    column = column_of(sheet, "月", "教員")
    assert sheet.cell(FIRST_BODY_ROW, column).value == "築　雅之"


def test_a_teacher_without_an_original_spelling_falls_back(tmp_path):
    """元表記が空でも氏名を落とさない。"""
    sheet = write(tmp_path, [make("A1", teacher="築雅之")],
                  {"A1": (TimeSlot("月", 1),)})["経営・前期"]
    assert sheet.cell(FIRST_BODY_ROW, column_of(sheet, "月", "教員")).value == "築雅之"


# ---------------------------------------------------------------- 横の区切り

def _bottom(sheet, row: int, column: int):
    """その行の下に引かれている線の種類。引かれていなければ None。

    **結合したセルでは見えない。** openpyxl は結合すると、その範囲の外枠を
    先頭セルに写す。科目名の列は教員の人数ぶん結合されるので、行と行の
    あいだを見るには結合していない列（教員）を見る。
    """
    side = sheet.cell(row, column).border.bottom
    return side.style if side else None


def test_the_rules_tell_periods_years_and_rows_apart(tmp_path):
    """横の区切りは 3 段階。**同じ年次の中には線を引かない。**

    1 科目が教員の人数ぶん行を使うので、1 行ごとに線を引くと科目の
    切れ目が分からなくなる。
    """
    # 1 年に 2 名、2 年に 1 名。1 時限だけを埋める。
    subjects = [
        make("A1", name="科目甲", year=1, teacher="教員1"),
        make("A2", name="科目甲", year=1, teacher="教員2"),
        make("A3", name="科目乙", year=2, teacher="教員3"),
    ]
    sheet = write(tmp_path, subjects,
                  {s.code: (TimeSlot("月", 1),) for s in subjects})["経営・前期"]

    teacher = column_of(sheet, "月", "教員")
    name = column_of(sheet, "月", "授業科目名")
    first = find_row(sheet, name, "科目甲")

    assert sheet.cell(first, teacher).value == "教員1"
    assert sheet.cell(first + 1, teacher).value == "教員2"
    assert _bottom(sheet, first, teacher) is None, "同じ年次の中には線を引かない"
    assert _bottom(sheet, first + 1, teacher) == "dotted", "年次のあいだは点線"

    # 1 時限の最後の年次（大学は 4 年）の下は太い実線になる
    last_of_period = find_row(sheet, 2, "4年")
    assert _bottom(sheet, last_of_period, teacher) == "medium", "時限のあいだは太い実線"
    assert _bottom(sheet, last_of_period, 1) == "medium"
    assert _bottom(sheet, last_of_period, 2) == "medium"


def test_the_last_period_is_closed_off(tmp_path):
    """5 時限の下も太い実線で閉じる。表の終わりが分かるようにする。"""
    sheet = write(tmp_path, [make("A1")], {"A1": (TimeSlot("月", 1),)})["経営・前期"]
    last = max(r for r in range(FIRST_BODY_ROW, sheet.max_row + 1)
               if sheet.cell(r, 2).value or _bottom(sheet, r, 2))
    assert _bottom(sheet, last, column_of(sheet, "月", "教員")) == "medium"
    assert _merged_label(sheet, last, 1) == "5時限"


# ---------------------------------------------------------------- 集中の終わり

def test_the_intensive_column_stops_where_the_subjects_run_out(tmp_path):
    """集中の欄は科目が尽きたところで太い実線を引いて終わる。

    グリッド側は 5 時限 × 年次ぶんの高さが必ず要るが、集中はあるだけしか
    無い。下まで空の枠を伸ばすと、何も無い場所を延々と目で追うことになる。
    """
    grid = make("A1", name="ふつうの科目")
    intensive = make("A9", name="認定ＰＢＬ", is_intensive=True, year=1)
    sheet = write(tmp_path, [grid, intensive], {"A1": (TimeSlot("月", 1),)},
                  intensive_codes=["A9"])["経営・前期"]

    year_column, name = intensive_block(sheet)
    row = find_row(sheet, name, "認定ＰＢＬ")
    assert _bottom(sheet, row, name) == "medium", "最後の科目の下は太い実線"

    # その下には枠を伸ばさない
    below = row + 1
    assert sheet.cell(below, name).border.left.style is None
    assert sheet.cell(below, year_column).fill.patternType is None
    assert sheet.max_row > below, "グリッド側はまだ続いている"


def test_a_one_line_name_shrinks_instead_of_wrapping(tmp_path):
    """Excel では wrap が shrink を打ち消す。

    1 行しかないセルで折り返すと、2 行目が行高に隠れて読めなくなる。
    縦に結合したセルは行数ぶんの高さがあるので、そこだけ折り返す。
    """
    alone = make("A1", name="ネットワークシステム開発実習", year=1)
    shared = [make(f"A{i}", name="課題研究Ⅰ", year=3, teacher=f"教員{i}")
              for i in range(2, 4)]
    sheet = write(tmp_path, [alone, *shared], {
        s.code: (TimeSlot("月", 1),) for s in [alone, *shared]
    })["経営・前期"]

    name = column_of(sheet, "月", "授業科目名")
    single = sheet.cell(find_row(sheet, name, "ネットワークシステム開発実習"), name)
    # False は保存時に省かれ、読み戻すと None になる
    assert not single.alignment.wrap_text
    assert single.alignment.shrink_to_fit is True

    merged = sheet.cell(find_row(sheet, name, "課題研究Ⅰ"), name)
    assert merged.alignment.wrap_text is True, "結合したセルは折り返してよい"
