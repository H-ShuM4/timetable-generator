"""カリキュラム一覧 Excel を Subject のリストに変換する。

▲科目は同一授業コードで 2 行に分かれているため 1 件へ集約する。
コース列・遠隔列は事務局側で整備中のため、存在しなくても落ちない。
"""
import json
from collections import defaultdict
from pathlib import Path

import openpyxl

from app.ingest.name_normalizer import normalize_name
from app.models.enums import Category, Department, Quarter, Term
from app.models.subject import Subject
from app.models.timeslot import INTENSIVE_PERIOD, TimeSlot

_CONFIG_DIR = Path(__file__).resolve().parents[2] / "config"

DOUBLE_SLOT_MARKER = "▲"
ACCOUNTING_SUFFIX = ":会"
RETAKE_MARKER = "【再】"
INTENSIVE_DAY = "集中"

COURSE_COLUMN_NAMES = ("コース", "フィールド")
"""コース列の見出し。大学は「コース」、短大は「フィールド」と呼ぶ。"""


def _load_json(filename: str, fallback):
    path = _CONFIG_DIR / filename
    if not path.exists():
        return fallback
    return json.loads(path.read_text(encoding="utf-8"))


def to_base_name(name: str) -> str:
    """先頭の ▲、接尾辞の :会、【再】を除いた正規化名を返す。"""
    text = str(name).strip()
    if text.startswith(DOUBLE_SLOT_MARKER):
        text = text[len(DOUBLE_SLOT_MARKER):]
    text = text.replace(ACCOUNTING_SUFFIX, "")
    text = text.replace(RETAKE_MARKER, "")
    return text.strip()


def _header_index(header_row: tuple) -> dict[str, int]:
    return {
        str(value).strip(): index
        for index, value in enumerate(header_row)
        if value is not None
    }


def _cell(row: tuple, columns: dict[str, int], key: str):
    index = columns.get(key)
    if index is None or index >= len(row):
        return None
    return row[index]


def _parse_quarter(value) -> Quarter | None:
    if not value:
        return None
    text = str(value).strip()
    try:
        return Quarter(text)
    except ValueError:
        return None


def _parse_slot(day_value, period_value) -> TimeSlot | None:
    """曜日と時限が揃っていて集中でない場合のみ TimeSlot を返す。"""
    if not day_value or period_value in (None, ""):
        return None
    day = str(day_value).strip()
    if day == INTENSIVE_DAY:
        return None
    period = int(period_value)
    if period == INTENSIVE_PERIOD:
        return None
    return TimeSlot(day, period)


def _is_intensive(day_value, period_value) -> bool:
    if day_value and str(day_value).strip() == INTENSIVE_DAY:
        return True
    if period_value in (None, ""):
        return False
    return int(period_value) == INTENSIVE_PERIOD


def _course_cell(row: tuple, columns: dict[str, int]):
    """シートごとに異なるコース列の見出しを吸収する。"""
    for name in COURSE_COLUMN_NAMES:
        if name in columns:
            return _cell(row, columns, name)
    return None


def _parse_courses(value) -> list[str]:
    if not value:
        return []
    return [token.strip() for token in str(value).replace("、", ",").split(",") if token.strip()]


def read_curriculum_rows(path: str | Path) -> list[dict]:
    """全シートの行を、列見出しをキーにした dict のリストとしてそのまま返す。

    Subject に変換すると曜日・時限が片方だけ入力されている行の情報は
    失われる（`fixed_slot` は両方揃った場合のみ作られる）ため、
    `validators.check_partial_slots` はこの生の行データを直接見る。
    """
    workbook = openpyxl.load_workbook(path, data_only=True)
    rows_out: list[dict] = []
    for sheet in workbook.worksheets:
        rows = list(sheet.iter_rows(values_only=True))
        if not rows:
            continue
        columns = _header_index(rows[0])
        for row in rows[1:]:
            code = _cell(row, columns, "授業コード")
            if not code:
                continue
            rows_out.append({name: _cell(row, columns, name) for name in columns})
    return rows_out


def read_curriculum(path: str | Path) -> list[Subject]:
    """全シートを読み、Subject のリストを返す。"""
    seminar_names = set(_load_json("seminar_subjects.json", []))
    overrides = _load_json("subject_overrides.json", {})
    non_consecutive = set(overrides.get("non_consecutive_double_subjects", []))

    workbook = openpyxl.load_workbook(path, data_only=True)
    grouped: dict[str, list[tuple[tuple, dict[str, int]]]] = defaultdict(list)
    order: list[str] = []

    for sheet in workbook.worksheets:
        rows = list(sheet.iter_rows(values_only=True))
        if not rows:
            continue
        columns = _header_index(rows[0])
        for row in rows[1:]:
            code = _cell(row, columns, "授業コード")
            if not code:
                continue
            code = str(code).strip()
            if code not in grouped:
                order.append(code)
            grouped[code].append((row, columns))

    subjects: list[Subject] = []
    for code in order:
        entries = grouped[code]
        row, columns = entries[0]
        name = str(_cell(row, columns, "授業科目名称")).strip()
        base_name = to_base_name(name)

        slots: list[TimeSlot] = []
        intensive = False
        for entry_row, entry_columns in entries:
            day = _cell(entry_row, entry_columns, "曜日")
            period = _cell(entry_row, entry_columns, "時限")
            if _is_intensive(day, period):
                intensive = True
                continue
            slot = _parse_slot(day, period)
            if slot is not None:
                slots.append(slot)

        is_double = name.startswith(DOUBLE_SLOT_MARKER)
        slots_required = 2 if is_double else 1
        requires_consecutive = is_double and base_name not in non_consecutive

        subjects.append(
            Subject(
                code=code,
                name=name,
                base_name=base_name,
                department=Department(str(_cell(row, columns, "学科")).strip()),
                year=int(str(_cell(row, columns, "年次配当")).strip()),
                term=Term(str(_cell(row, columns, "開講期間")).strip()),
                quarter=_parse_quarter(_cell(row, columns, "備考")),
                category=Category(str(_cell(row, columns, "科目区分")).strip()),
                courses=_parse_courses(_course_cell(row, columns)),
                teacher=normalize_name(_cell(row, columns, "教員氏名")),
                is_remote=str(_cell(row, columns, "遠隔") or "").strip() == "○",
                is_joint=str(_cell(row, columns, "合同(経・会)") or "").strip() == "○",
                slots_required=slots_required,
                requires_consecutive=requires_consecutive,
                fixed_slot=tuple(slots) if len(slots) == slots_required else None,
                is_intensive=intensive,
                is_seminar=base_name in seminar_names,
            )
        )

    return subjects
